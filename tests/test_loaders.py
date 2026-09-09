import io

import pytest
from docx import Document as DocxDocument
from pypdf import PdfWriter

from app.services.ingestion.loaders import (
    LoaderError,
    load_document,
    load_docx,
    load_markdown,
    load_pdf,
    load_text,
)


def _blank_pdf(pages: int = 2) -> bytes:
    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def _docx_bytes() -> bytes:
    doc = DocxDocument()
    doc.add_heading("Leave Policy", level=1)
    doc.add_paragraph("Employees receive 20 days of annual leave.")
    doc.add_heading("Contractors", level=1)
    doc.add_paragraph("Contractors accrue leave pro rata.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_pdf_reports_page_count_and_no_text_for_blank_pages():
    blocks, page_count = load_pdf(_blank_pdf(3))
    assert page_count == 3
    assert blocks == []


def test_pdf_rejects_garbage():
    with pytest.raises(LoaderError):
        load_pdf(b"definitely not a pdf")


def test_docx_extracts_sections():
    blocks, page_count = load_docx(_docx_bytes())
    assert page_count is None
    sections = [b.section for b in blocks]
    assert "Leave Policy" in sections
    assert "Contractors" in sections
    assert any("20 days" in b.text for b in blocks)


def test_docx_rejects_garbage():
    with pytest.raises(LoaderError):
        load_docx(b"not a docx")


def test_text_splits_on_blank_lines():
    blocks, page_count = load_text(b"First paragraph.\n\nSecond paragraph.")
    assert page_count is None
    assert [b.text for b in blocks] == ["First paragraph.", "Second paragraph."]


def test_markdown_keeps_headings_as_sections():
    md = b"# Annual Leave\n\nEmployees get 20 days.\n\n## Carryover\n\nUp to 5 days."
    blocks, _ = load_markdown(md)
    by_section = {b.section: b.text for b in blocks}
    assert "Annual Leave" in by_section
    assert "Carryover" in by_section
    assert "20 days" in by_section["Annual Leave"]


def test_load_document_dispatches_and_rejects_unknown():
    blocks, _ = load_document("txt", b"hello world")
    assert blocks[0].text == "hello world"
    with pytest.raises(LoaderError):
        load_document("rtf", b"data")
